"""PPTX 일반 문서의 슬라이드·표·Markdown 변환 회귀 테스트."""
from pathlib import Path
from types import SimpleNamespace

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import service.parsers as parsers_module
from service.parsers import parse_pptx
from service.worker import _parse


def _create_sample_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "AWS Bedrock 구성"

    textbox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(6), Inches(1))
    textbox.text_frame.text = "A서버에서 Bedrock endpoint를 설정합니다."

    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1.5)).table
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    table.cell(1, 0).text = "Region"
    table.cell(1, 1).text = "ap-northeast-2"

    image_path = path.with_suffix(".png")
    Image.new("RGB", (320, 120), color="white").save(image_path)
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(5), Inches(4), Inches(1.5))
    presentation.save(path)


def _settings(ocr: str = "disabled", device: str = "cpu") -> SimpleNamespace:
    return SimpleNamespace(
        solutions=SimpleNamespace(
            ocr=ocr,
            easyocr_device=device,
            easyocr_languages=["ko", "en"],
            docling=SimpleNamespace(
                artifacts_path="",
                require_local_artifacts=False,
            ),
        )
    )


def test_parse_pptx_extracts_slide_text_table_and_markdown(monkeypatch, tmp_path):
    monkeypatch.setattr(parsers_module, "get_settings", lambda: _settings())
    source = tmp_path / "guide.pptx"
    output = tmp_path / "processed"
    _create_sample_pptx(source)

    items, report = parse_pptx(source, output)

    assert report["slides"] == 1
    assert report["tables"] == 1
    assert report["images"] == 1
    assert report["ocr_images"] == 0
    assert any("A서버에서 Bedrock endpoint" in item.text for item in items)
    assert any(
        "항목: Region" in item.text and "값: ap-northeast-2" in item.text
        for item in items
    )
    assert all(item.page_start == 1 and item.page_end == 1 for item in items)
    markdown = (output / "document.md").read_text(encoding="utf-8")
    assert "## 1. AWS Bedrock 구성" in markdown
    assert "| 항목 | 값 |" in markdown
    assert (output / "chunks.jsonl").is_file()
    assert (output / "metadata.json").is_file()


def test_parse_pptx_uses_easyocr_for_picture_text(monkeypatch, tmp_path):
    class FakeReader:
        def readtext(self, image, detail, paragraph):
            assert image.shape[0] > 0 and image.shape[1] > 0
            assert detail == 0
            assert paragraph is False
            return ["이미지 내부 설정값", "endpoint 10.2.12.98"]

    monkeypatch.setattr(
        parsers_module,
        "get_settings",
        lambda: _settings(ocr="easyocr", device="cpu"),
    )
    monkeypatch.setattr(
        parsers_module,
        "_get_easyocr_reader",
        lambda *args: FakeReader(),
    )
    source = tmp_path / "ocr-guide.pptx"
    output = tmp_path / "processed"
    _create_sample_pptx(source)

    items, report = parse_pptx(source, output)

    ocr_items = [
        item for item in items if item.metadata.get("content_source") == "image_ocr"
    ]
    assert len(ocr_items) == 1
    assert "이미지 내부 설정값" in ocr_items[0].text
    assert "endpoint 10.2.12.98" in ocr_items[0].text
    assert ocr_items[0].page_start == 1
    assert report["ocr_images"] == 1
    assert report["ocr_device"] == "cpu"
    assert "### 이미지 OCR" in (output / "document.md").read_text(encoding="utf-8")


def test_worker_dispatches_pptx_to_dedicated_parser(monkeypatch, tmp_path):
    expected = ([], {"slides": 0})
    calls = []

    def fake_parser(source, output):
        calls.append((source, output))
        return expected

    monkeypatch.setattr("service.worker.parse_pptx", fake_parser)
    source = tmp_path / "guide.pptx"
    output = tmp_path / "processed"

    assert _parse({"asset_type": "document"}, source, output) == expected
    assert calls == [(source, output)]
